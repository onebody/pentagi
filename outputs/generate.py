"""PentAGI 项目演示 PPT 生成脚本（python-pptx）。

用法: python generate.py
输出: outputs/PentAGI-项目演示.pptx
"""

import math
import os
import re
import zipfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

PALETTE = {
    "dominant": RGBColor(0x0E, 0x22, 0x33),   # deep slate navy
    "support1": RGBColor(0xA8, 0xC6, 0xDA),   # dusty blue
    "support2": RGBColor(0xF2, 0xF5, 0xF7),   # off-white card
    "accent": RGBColor(0x3B, 0xD4, 0x8A),     # terminal green
    "alert": RGBColor(0xF2, 0x65, 0x5A),      # coral red (risk only)
    "ink": RGBColor(0x16, 0x20, 0x2A),
    "muted": RGBColor(0x66, 0x76, 0x84),
    "border": RGBColor(0xD8, 0xE2, 0xEA),
    "white": RGBColor(0xFF, 0xFF, 0xFF),
}

FONT_HEAD = "Arial"
FONT_BODY = "Arial"
FONT_MONO = "Courier New"
FONT_EA = "PingFang SC"

WARNINGS = []

# ---------- text width estimator (mixed CJK / latin) ----------
def _char_w(ch, mono=False):
    o = ord(ch)
    if mono:
        return 0.60
    if o >= 0x2E80:
        return 1.0
    if ch == " ":
        return 0.30
    if ch.isdigit():
        return 0.56
    if ch.isupper():
        return 0.66
    if ch in ".,:;'|!()[]/\\-_=+·":
        return 0.40
    return 0.52


def est_lines(text, width_in, size_pt, mono=False):
    usable = max(width_in - 0.12, 0.5)
    total = sum(_char_w(c, mono) for c in text) * size_pt / 72.0
    return max(1, math.ceil(total / usable * 1.06))


def line_h(size_pt):
    return size_pt * 1.34 / 72.0


def guard(text, key, max_chars):
    if len(text) <= max_chars:
        return text
    return text[: max_chars - 1] + "…"


# ---------- primitives ----------
def add_text(slide, text, x, y, w, h, *, size=14, bold=False, color=None,
             align=PP_ALIGN.LEFT, mono=False, name=None, tag="", slack=0.08):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = name or (FONT_MONO if mono else FONT_BODY)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color if color is not None else PALETTE["ink"]
    rPr = r._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {"typeface": FONT_EA})
        rPr.append(ea)
    else:
        ea.set("typeface", FONT_EA)

    need = est_lines(text, w, size, mono) * line_h(size) + slack
    if need > h + 0.005:
        WARNINGS.append(f"[{tag}] overflow: need {need:.2f}in > {h:.2f}in :: {text[:34]}")
    return tb


def add_multiline(slide, lines, x, y, w, h, *, size=14, color=None,
                  bold=False, mono=False, gap=0.06, tag=""):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    total = 0.0
    for i, txt in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = txt
        r.font.name = FONT_MONO if mono else FONT_BODY
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color if color is not None else PALETTE["ink"]
        rPr = r._r.get_or_add_rPr()
        ea = rPr.find(qn("a:ea"))
        if ea is None:
            ea = rPr.makeelement(qn("a:ea"), {"typeface": FONT_EA})
            rPr.append(ea)
        else:
            ea.set("typeface", FONT_EA)
        total += est_lines(txt, w, size, mono) * line_h(size) + gap
    if total > h + 0.02:
        WARNINGS.append(f"[{tag}] block overflow: need {total:.2f}in > {h:.2f}in")
    return tb


def add_rect(slide, x, y, w, h, fill, *, line=None, rounded=False):
    shape = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    return shp


def card(slide, x, y, w, h, *, stripe=None):
    add_rect(slide, x, y, w, h, PALETTE["support2"], line=PALETTE["border"], rounded=True)
    if stripe:
        add_rect(slide, x, y + 0.12, 0.09, h - 0.24, stripe)


def chip(slide, text, x, y, w, h, *, fill, color, size=13, mono=False, bold=False):
    add_rect(slide, x, y, w, h, fill, rounded=True)
    add_text(slide, text, x, y + (h - line_h(size)) / 2 - 0.02, w, h,
             size=size, color=color, align=PP_ALIGN.CENTER, mono=mono,
             bold=bold, tag="chip")


def chrome(slide, title, kicker, page, total):
    add_rect(slide, 0, 0, 0.16, 7.5, PALETTE["accent"])
    add_text(slide, kicker, 0.62, 0.42, 11.9, 0.35, size=12, bold=True,
             color=PALETTE["accent"], tag="kicker")
    add_text(slide, title, 0.6, 0.78, 11.9, 0.85, size=29, bold=True,
             color=PALETTE["dominant"], name=FONT_HEAD, tag="title")
    add_text(slide, "PentAGI 项目演示", 0.62, 7.06, 6.0, 0.32, size=10,
             color=PALETTE["muted"], tag="foot")
    add_text(slide, f"{page} / {total}", 11.3, 7.06, 1.4, 0.32, size=10,
             color=PALETTE["muted"], align=PP_ALIGN.RIGHT, tag="page")


TOTAL = 10
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def new():
    return prs.slides.add_slide(blank)


# ---------- S1 cover ----------
s = new()
add_rect(s, 0, 0, 13.333, 7.5, PALETTE["dominant"])
add_rect(s, 0, 0, 0.22, 7.5, PALETTE["accent"])
add_text(s, "SELF-HOSTED  ·  AI  PENETRATION TESTING ORCHESTRATION", 0.9, 1.55, 11.5, 0.4,
         size=13, bold=True, color=PALETTE["accent"], mono=True, tag="cover-kicker")
add_text(s, "PentAGI 项目演示", 0.85, 2.05, 11.6, 1.3, size=52, bold=True,
         color=PALETTE["white"], name=FONT_HEAD, tag="cover-title")
add_text(s, "自主渗透测试多 Agent 平台 —— 架构、能力与交付现状", 0.9, 3.45, 11.4, 0.7,
         size=21, color=PALETTE["support1"], tag="cover-sub")
add_multiline(s, [
    "$ docker compose up -d",
    "$ open https://localhost:8443",
], 0.9, 4.45, 5.3, 1.15, size=14, color=PALETTE["accent"], mono=True, gap=0.10, tag="term")
add_rect(s, 6.6, 4.45, 5.85, 1.15, None, line=PALETTE["support1"])
add_multiline(s, [
    "上游  github.com/vxcontrol/pentagi   (MIT)",
    "镜像  github.com/onebody/pentagi     main",
], 6.85, 4.68, 5.4, 0.85, size=13, color=PALETTE["support1"], mono=True, gap=0.10, tag="repo")
add_text(s, "2026-10-14  ·  基于仓库快照与 docs/01–03 中文文档整理", 0.9, 6.85, 11.5, 0.4,
         size=12, color=PALETTE["support1"], tag="cover-foot")

# ---------- S2 定位与核心能力 ----------
s = new()
chrome(s, "它解决什么问题", "SYSTEM POSITIONING", 2, TOTAL)
add_text(s, "把一次渗透测试委托给 AI：平台负责编排模型、工具、沙箱与证据留存",
         0.62, 1.72, 11.9, 0.5, size=16, color=PALETTE["muted"], tag="s2-lead")
cards = [
    ("Flow 为中心的任务单元", "目标、上下文、执行过程、产物与报告收进同一个容器，可复盘可追溯"),
    ("多 Agent + 多工具编排", "Pentester / Coder / Searcher 等角色轮转，调用终端、浏览器、搜索、向量库"),
    ("Docker 沙箱执行", "命令与工具在 Flow 专属任务容器内运行，宿主环境隔离"),
    ("统一控制面", "Web UI + REST + GraphQL(WS)，管理 Flow / 模板 / Provider / 提示词 / Token"),
]
w, h, gx, gy = 5.95, 2.05, 0.45, 0.35
for i, (head, body) in enumerate(cards):
    x = 0.62 + (i % 2) * (w + gx)
    y = 2.45 + (i // 2) * (h + gy)
    card(s, x, y, w, h, stripe=PALETTE["accent"])
    add_rect(s, x + 0.28, y + 0.32, 0.42, 0.42, PALETTE["dominant"])
    add_text(s, str(i + 1), x + 0.28, y + 0.36, 0.42, 0.38, size=15, bold=True,
             color=PALETTE["white"], align=PP_ALIGN.CENTER, tag="num")
    add_text(s, head, x + 0.88, y + 0.32, w - 1.15, 0.45, size=18, bold=True,
             color=PALETTE["dominant"], name=FONT_HEAD, tag="s2-head")
    add_text(s, body, x + 0.88, y + 0.95, w - 1.15, 0.95, size=14,
             color=PALETTE["ink"], tag="s2-body")

# ---------- S3 规模数字 ----------
s = new()
chrome(s, "仓库与代码规模", "BY THE NUMBERS", 3, TOTAL)
stats = [
    ("25.6 万行", "后端 Go  ·  518 个 .go 文件"),
    ("3.5 万行", "前端 TS/TSX  ·  157 个源文件"),
    ("71 MB", "仓库体积 · 含 54 MB jaeger-clickhouse 二进制"),
    ("14 / 19", "Provider 配置模板 / 同题横评报告"),
]
w, h, gx, gy = 5.95, 2.0, 0.45, 0.4
for i, (num, label) in enumerate(stats):
    x = 0.62 + (i % 2) * (w + gx)
    y = 1.95 + (i // 2) * (h + gy)
    card(s, x, y, w, h)
    add_text(s, num, x + 0.3, y + 0.28, w - 0.6, 1.0, size=44, bold=True,
             color=PALETTE["dominant"], align=PP_ALIGN.CENTER, name=FONT_HEAD, tag="s3-num")
    add_text(s, label, x + 0.3, y + 1.35, w - 0.6, 0.5, size=14,
             color=PALETTE["muted"], align=PP_ALIGN.CENTER, tag="s3-label")
add_rect(s, 0.62, 6.28, 12.05, 0.62, PALETTE["support2"], line=PALETTE["border"], rounded=True)
add_text(s, "另有 4 套 Compose 编排栈、622 行中文文档（docs/01–03）与 4 份 backend/docs 二开专题",
         0.9, 6.4, 11.6, 0.4, size=13, color=PALETTE["ink"], tag="s3-foot")

# ---------- S4 架构分层 ----------
s = new()
chrome(s, "五层架构", "ARCHITECTURE", 4, TOTAL)
layers = [
    ("接入层", "React + TS + Vite + Apollo SPA   ·   Gin REST /api/v1   ·   gqlgen GraphQL + WS subscriptions"),
    ("编排层", "FlowController → FlowWorker → Task / SubTask 状态机，驱动多 Agent 轮转与工具调用"),
    ("能力层", "工具执行器：Terminal   Scraper 浏览器   Web Search   Vector Store   Graphiti 知识图谱"),
    ("模型层", "Provider 适配包：openai · anthropic · gemini · bedrock · ollama · qwen · glm · kimi · deepseek · custom + embeddings"),
    ("数据与观测", "PostgreSQL + pgvector（goose 迁移 / sqlc）· 截图与产物存储 · 可选 Langfuse / Grafana+Loki+Jaeger / Neo4j"),
]
y = 1.85
for name, desc in layers:
    add_rect(s, 0.62, y, 2.35, 0.86, PALETTE["dominant"], rounded=False)
    add_text(s, name, 0.62, y + 0.28, 2.35, 0.45, size=16, bold=True,
             color=PALETTE["white"], align=PP_ALIGN.CENTER, name=FONT_HEAD, tag="s4-name")
    add_rect(s, 3.12, y, 9.55, 0.86, PALETTE["support2"], line=PALETTE["border"])
    add_text(s, desc, 3.38, y + 0.16, 9.05, 0.6, size=13.5,
             color=PALETTE["ink"], tag="s4-desc")
    y += 1.0
add_rect(s, 0.62, y + 0.02, 12.05, 0.72, PALETTE["dominant"], rounded=True)
add_text(s, "执行底座：后端通过 Docker client 为每个 Flow 创建沙箱容器，所有命令在容器内执行并回传日志",
         0.92, y + 0.22, 11.5, 0.42, size=14, bold=True, color=PALETTE["accent"],
         mono=False, tag="s4-bar")

# ---------- S5 执行链路 ----------
s = new()
chrome(s, "一次 Flow 的执行链路", "RUNTIME FLOW", 5, TOTAL)
steps = [
    ("创建 Flow", "绑定目标、语言、模板"),
    ("装配 Worker", "加载 Provider 与工具"),
    ("下发指令", "范围 / 禁止动作 / 输出要求"),
    ("Agent 推理", "规划 Task 与 SubTask"),
    ("沙箱执行", "终端、浏览器、搜索调用"),
    ("证据与报告", "日志 / 截图 / 报告落库"),
]
w, gap = 1.9, 0.11
for i, (label, sub) in enumerate(steps):
    x = 0.62 + i * (w + gap)
    add_rect(s, x, 2.15, w, 0.62, PALETTE["dominant"])
    add_text(s, f"0{i+1}", x, 2.28, w, 0.4, size=16, bold=True,
             color=PALETTE["accent"], align=PP_ALIGN.CENTER, mono=True, tag="s5-idx")
    add_rect(s, x, 2.85, w, 1.5, PALETTE["support2"], line=PALETTE["border"])
    add_text(s, label, x + 0.08, 3.0, w - 0.16, 0.5, size=15, bold=True,
             color=PALETTE["dominant"], align=PP_ALIGN.CENTER, name=FONT_HEAD, tag="s5-label")
    add_text(s, sub, x + 0.12, 3.52, w - 0.24, 0.75, size=12,
             color=PALETTE["muted"], align=PP_ALIGN.CENTER, tag="s5-sub")
    if i < len(steps) - 1:
        add_text(s, "▸", x + w - 0.02, 2.2, 0.2, 0.5, size=16, bold=True,
                 color=PALETTE["accent"], align=PP_ALIGN.CENTER, tag="arrow")
add_rect(s, 0.62, 4.72, 12.05, 0.78, PALETTE["support2"], line=PALETTE["border"], rounded=True)
add_text(s, "全程通过 GraphQL subscriptions 实时推送：助手/自动化消息、工具调用日志、终端输出、截图与向量库读写",
         0.92, 4.98, 11.45, 0.42, size=14, color=PALETTE["ink"], tag="s5-note")
add_text(s, "对象模型：Flow → Task → SubTask，Container 承载执行环境，Logs / Screenshot / Prompt / Template 构成审计面",
         0.62, 5.85, 12.05, 0.7, size=14, color=PALETTE["muted"], tag="s5-obj")

# ---------- S6 Agent 矩阵 ----------
s = new()
chrome(s, "Agent 角色分工", "AGENT ROLES", 6, TOTAL)
roles = [
    ("Pentester", "核心执行角色：规划攻击面、下达并验证动作"),
    ("Coder", "生成脚本 / Payload / 工具代码并写入容器"),
    ("Searcher", "Web 搜索与知识检索，补齐上下文与证据"),
]
y = 1.9
for name, desc in roles:
    card(s, 0.62, y, 5.75, 1.42, stripe=PALETTE["accent"])
    add_text(s, name, 0.95, y + 0.2, 5.1, 0.45, size=19, bold=True,
             color=PALETTE["dominant"], name=FONT_MONO, tag="s6-role")
    add_text(s, desc, 0.95, y + 0.72, 5.1, 0.6, size=13.5,
             color=PALETTE["ink"], tag="s6-desc")
    y += 1.62
add_rect(s, 6.62, 1.9, 6.05, 4.46, PALETTE["support2"], line=PALETTE["border"])
add_text(s, "同一框架内的其他 AgentType（源码节选）", 6.9, 2.1, 5.5, 0.45, size=15, bold=True,
         color=PALETTE["dominant"], name=FONT_HEAD, tag="s6-side-title")
others = ["Assistant", "Reflector", "Reporter", "Adviser", "Memorist", "Enricher",
          "Summarizer", "Refiner", "Generator", "ToolCallFixer", "Installer", "UsageStats"]
cw, ch, gx, gy2 = 1.72, 0.5, 0.12, 0.14
for i, name in enumerate(others):
    x = 6.9 + (i % 3) * (cw + gx)
    yy = 2.72 + (i // 3) * (ch + gy2)
    chip(s, name, x, yy, cw, ch, fill=PALETTE["white"], color=PALETTE["ink"],
         size=12, mono=True)
add_text(s, "CLAUDE.md 对外把这套角色简化表述为 Researcher / Developer / Executor",
         6.9, 5.6, 5.5, 0.62, size=13, color=PALETTE["muted"], tag="s6-note")

# ---------- S7 Provider 生态 ----------
s = new()
chrome(s, "模型 Provider 生态", "LLM PROVIDERS", 7, TOTAL)
providers = [
    "custom-openai", "deepinfra", "deepseek", "moonshot",
    "novita", "openrouter", "ollama-cloud", "ollama-llama318b",
    "ollama-llama318b-instruct", "ollama-qwen332b-fp16-tc", "ollama-qwq32b-fp16-tc",
    "vllm-qwen3.5-27b-fp8", "vllm-qwen3.5-27b-fp8-no-think", "vllm-qwen332b-fp16",
    "共 14 份模板",
]
cw, ch, gx, gy3 = 2.85, 0.56, 0.22, 0.16
for i, name in enumerate(providers):
    x = 0.62 + (i % 4) * (cw + gx)
    y = 1.95 + (i // 4) * (ch + gy3)
    last = i == len(providers) - 1
    chip(s, name, x, y, cw, ch,
         fill=PALETTE["dominant"] if last else PALETTE["support2"],
         color=PALETTE["white"] if last else PALETTE["ink"],
         size=11.5 if (name.startswith("vllm") or name.startswith("ollama-")) and not last else 13,
         mono=not last)
add_rect(s, 0.62, 4.95, 12.05, 1.75, PALETTE["dominant"], rounded=True)
add_multiline(s, [
    "代码内置 10 类适配包 + embeddings；OpenRouter / Novita / DeepInfra / vLLM 走 OpenAI 兼容 custom 路径",
    "examples/tests 有 19 份同题横评报告（base_web_pentest），可直接对比换模型对渗透效果的影响",
    "examples/guides 另附 vLLM-Qwen3.5-27B 与 worker node 两份私有化部署指南",
], 0.95, 5.22, 11.4, 1.3, size=14, color=PALETTE["white"], gap=0.12, tag="s7-note")

# ---------- S8 部署形态 ----------
s = new()
chrome(s, "部署形态与运行基线", "DEPLOYMENT", 8, TOTAL)
stacks = [
    ("docker-compose.yml", "pentagi · pgvector · pgexporter · scraper（必需）"),
    ("+ docker-compose-observability.yml", "Grafana · OTel · Loki · Jaeger-ClickHouse · VictoriaMetrics"),
    ("+ docker-compose-langfuse.yml", "Langfuse：LLM 调用观测与评估"),
    ("+ docker-compose-graphiti.yml", "Neo4j + Graphiti：知识图谱长期记忆"),
]
y = 1.9
for name, desc in stacks:
    add_rect(s, 0.62, y, 6.35, 1.02, PALETTE["support2"], line=PALETTE["border"])
    add_text(s, name, 0.85, y + 0.14, 6.0, 0.4, size=14, bold=True,
             color=PALETTE["dominant"], mono=True, tag="s8-name")
    add_text(s, desc, 0.85, y + 0.55, 6.0, 0.4, size=12.5,
             color=PALETTE["muted"], tag="s8-desc")
    y += 1.18
card(s, 7.32, 1.9, 5.35, 4.42, stripe=PALETTE["alert"])
add_text(s, "上线前基线", 7.62, 2.1, 4.8, 0.45, size=17, bold=True,
         color=PALETTE["dominant"], name=FONT_HEAD, tag="s8-card-title")
add_multiline(s, [
    "· 默认仅绑定 127.0.0.1，入口 https://localhost:8443",
    "· 必须改：COOKIE_SIGNING_SALT、PUBLIC_URL、CORS_ORIGINS",
    "· 对外提供服务用正式 TLS，关闭 Swagger / Playground",
    "· 硬件：4C/8G 起，16G+ 推荐，磁盘 20G+ 且持续增长",
    "· 依赖 Docker Engine + Compose v2，需挂载 docker.sock",
], 7.62, 2.68, 4.85, 3.5, size=13.5, color=PALETTE["ink"], gap=0.14, tag="s8-card")

# ---------- S9 交付状态 ----------
s = new()
chrome(s, "本次交付：文档与版本", "DELIVERY STATUS", 9, TOTAL)
add_rect(s, 0.62, 1.85, 6.0, 0.62, PALETTE["accent"])
add_text(s, "已入库  main · f0e6a91", 0.62, 2.0, 6.0, 0.4, size=16, bold=True,
         color=PALETTE["dominant"], align=PP_ALIGN.CENTER, name=FONT_HEAD, tag="s9-h1")
add_multiline(s, [
    "docs/01 功能清单与操作手册 · 160 行",
    "docs/02 部署与运维文档 · 238 行",
    "25 处 file:/// 绝对链接改写为仓库相对路径",
    "推送前逐条校验链接在仓库内均可解析",
], 0.85, 2.68, 5.6, 2.6, size=14, color=PALETTE["ink"], gap=0.16, tag="s9-left")
add_rect(s, 6.9, 1.85, 5.77, 0.62, PALETTE["dominant"])
add_text(s, "仅本地  不进公开仓库", 6.9, 2.0, 5.77, 0.4, size=16, bold=True,
         color=PALETTE["white"], align=PP_ALIGN.CENTER, name=FONT_HEAD, tag="s9-h2")
add_multiline(s, [
    "docs/03 渗透测试报告 · 224 行 · 8 条 finding",
    "最高等级 Critical，含配置与代码定位",
    "针对活跃的上游项目：公开即等于漏洞披露",
    "结论：本次推送按约定剔除，明细留在本地文档",
], 7.15, 2.68, 5.4, 2.6, size=14, color=PALETTE["ink"], gap=0.16, tag="s9-right")
add_rect(s, 0.62, 5.55, 12.05, 1.15, PALETTE["support2"], line=PALETTE["border"], rounded=True)
add_multiline(s, [
    "镜像关系：onebody/pentagi 是 vxcontrol/pentagi 的 MIT 快照，LICENSE / NOTICE / EULA 均随仓保留",
    "本次提交只新增 2 个文件，未触碰任何上游代码；与远端 HEAD 逐文件比对无差异",
], 0.92, 5.82, 11.5, 0.85, size=13.5, color=PALETTE["ink"], gap=0.14, tag="s9-foot")

# ---------- S10 下一步 ----------
s = new()
add_rect(s, 0, 0, 13.333, 7.5, PALETTE["dominant"])
add_rect(s, 0, 0, 0.22, 7.5, PALETTE["accent"])
add_text(s, "下一步", 0.95, 1.35, 11.4, 0.9, size=38, bold=True,
         color=PALETTE["white"], name=FONT_HEAD, tag="s10-title")
nexts = [
    ("01", "把本地目录变成正式工作副本", "当前快照无 .git，且 .gitignore / .env.example / .github 等点文件在拷贝时丢失"),
    ("02", "为 docs/03 建立长期隔离", "临时 exclude 随 scratch clone 清理失效，改为持久忽略或独立私有仓"),
    ("03", "沉淀 Provider 横评与模板库", "把 19 份报告扩展成新模型对比基线，产出可复用的 Flow 模板与提示词"),
]
y = 2.55
for idx, head, sub in nexts:
    add_text(s, idx, 0.95, y, 0.9, 0.5, size=22, bold=True,
             color=PALETTE["accent"], mono=True, tag="s10-idx")
    add_text(s, head, 2.0, y, 5.2, 0.5, size=19, bold=True,
             color=PALETTE["white"], name=FONT_HEAD, tag="s10-h")
    add_text(s, sub, 7.35, y + 0.05, 5.0, 1.1, size=13.5,
             color=PALETTE["support1"], tag="s10-s")
    y += 1.28
add_text(s, "github.com/onebody/pentagi   ·   main @ f0e6a91   ·   2026-10-14",
         0.95, 6.85, 11.4, 0.4, size=12, color=PALETTE["support1"], mono=True, tag="s10-foot")

# ---------- save ----------
def save_pptx(prs, path):
    prs.save(path)
    tmp = path + ".tmp"
    try:
        with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(tmp, "w") as zout:
            for item in zin.infolist():
                if item.filename.lower().startswith("docprops/thumbnail"):
                    continue
                data = zin.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    data = re.sub(
                        r'<Override[^>]*PartName="/docProps/thumbnail[^"]*"[^>]*/?>',
                        '', data.decode("utf-8"), flags=re.IGNORECASE).encode("utf-8")
                if item.filename == "_rels/.rels":
                    data = re.sub(
                        r'<Relationship[^>]*Target="docProps/thumbnail[^"]*"[^>]*/?>',
                        '', data.decode("utf-8"), flags=re.IGNORECASE).encode("utf-8")
                zout.writestr(item, data)
        os.replace(tmp, path)
    except Exception as exc:
        print("warning: failed to strip template thumbnail:", exc)
        if os.path.exists(tmp):
            os.remove(tmp)


here = os.path.dirname(os.path.abspath(__file__))
out = os.path.join(here, "PentAGI-项目演示.pptx")
save_pptx(prs, out)
print("wrote", out, f"{os.path.getsize(out) / 1024:.0f} KB")
if WARNINGS:
    print(f"\n!! {len(WARNINGS)} layout warning(s):")
    for w_ in WARNINGS:
        print("  -", w_)
else:
    print("no overflow warnings")
